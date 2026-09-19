import sys
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(val='single', sz='4', color='E5E7EB', space='0')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b_xml = f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", "4")}" w:space="0" w:color="{edge_data.get("color", "E5E7EB")}"/>'
            tcBorders.append(parse_xml(b_xml))
    tcPr.append(tcBorders)

def add_styled_paragraph(doc, text, style='Normal', space_before=0, space_after=6, line_spacing=1.15):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    return p

def format_inlines(p, text, base_font="Calibri", base_size=10.5, base_color="374151"):
    # Tokenize bold and inline code: **bold** or `code` or normal
    # We can use regex to split tokens
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        run = p.add_run()
        run.font.name = base_font
        if token.startswith('**') and token.endswith('**'):
            run.text = token[2:-2]
            run.bold = True
            run.font.size = Pt(base_size)
            r, g, b = int(base_color[0:2], 16), int(base_color[2:4], 16), int(base_color[4:6], 16)
            run.font.color.rgb = RGBColor(r, g, b)
        elif token.startswith('`') and token.endswith('`'):
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 1)
            run.font.color.rgb = RGBColor(180, 40, 75)
            # soft highlight or dark pink
        else:
            run.text = token
            run.font.size = Pt(base_size)
            r, g, b = int(base_color[0:2], 16), int(base_color[2:4], 16), int(base_color[4:6], 16)
            run.font.color.rgb = RGBColor(r, g, b)

def build_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # Page setup (A4, 0.75 in margins)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    # Primary colors
    PRIMARY_COLOR = RGBColor(225, 29, 94)    # Pink / Rose #E11D5E
    SECONDARY_COLOR = RGBColor(234, 88, 12)  # Saffron/Orange #EA580C
    DARK_TEXT = RGBColor(30, 41, 59)         # Slate 800
    MUTED_TEXT = RGBColor(100, 116, 139)     # Slate 500

    i = 0
    total_lines = len(lines)

    while i < total_lines:
        line = lines[i].rstrip('\r\n')
        
        # Empty line
        if not line.strip():
            i += 1
            continue

        # Header 1: # Title
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[2:].strip())
            run.font.name = "Arial"
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = PRIMARY_COLOR
            i += 1
            continue

        # Subtitle or bold tagline immediately below H1
        if line.startswith('**Women-First Travel'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(line.replace('*', '').strip())
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = SECONDARY_COLOR
            i += 1
            continue

        # Header 2: ## Section
        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[3:].strip())
            run.font.name = "Arial"
            run.font.size = Pt(14)
            run.bold = True
            run.font.color.rgb = PRIMARY_COLOR

            # Add a stylish colored bottom accent border / line under H2
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="4" w:color="F43F5E"/></w:pBdr>')
            pPr.append(pBdr)

            i += 1
            continue

        # Header 3: ### Subsection
        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[4:].strip())
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = SECONDARY_COLOR
            i += 1
            continue

        # Header 4: #### Sub-subsection
        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[5:].strip())
            run.font.name = "Arial"
            run.font.size = Pt(11)
            run.bold = True
            run.font.color.rgb = DARK_TEXT
            i += 1
            continue

        # Horizontal rule ---
        if line.strip() == '---':
            # Subtle separator or empty paragraph with top border
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="1" w:color="E2E8F0"/></w:pBdr>')
            pPr.append(pBdr)
            i += 1
            continue

        # Blockquote: > text
        if line.startswith('> '):
            blockquote_lines = []
            while i < total_lines and lines[i].startswith('> '):
                blockquote_lines.append(lines[i][2:].rstrip('\r\n'))
                i += 1
            
            # Put in callout box table
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            cell.width = Inches(6.77)
            set_cell_background(cell, "FFF5F5")  # Soft rose/pink tint
            set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
            set_cell_border(cell, left=dict(val='single', sz='24', color='F43F5E', space='0'),
                                  top=dict(val='none'), bottom=dict(val='none'), right=dict(val='none'))
            
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            cp.paragraph_format.line_spacing = 1.15
            for idx, bl in enumerate(blockquote_lines):
                if idx > 0:
                    cp = cell.add_paragraph()
                    cp.paragraph_format.space_before = Pt(2)
                    cp.paragraph_format.space_after = Pt(2)
                    cp.paragraph_format.line_spacing = 1.15
                format_inlines(cp, bl, base_font="Arial", base_size=9.5, base_color="881337")
            
            # Add small gap after callout
            post_p = doc.add_paragraph()
            post_p.paragraph_format.space_before = Pt(0)
            post_p.paragraph_format.space_after = Pt(4)
            continue

        # Code block: ```
        if line.startswith('```'):
            code_lines = []
            i += 1
            while i < total_lines and not lines[i].startswith('```'):
                code_lines.append(lines[i].rstrip('\r\n'))
                i += 1
            if i < total_lines and lines[i].startswith('```'):
                i += 1 # skip closing ```

            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            cell.width = Inches(6.77)
            set_cell_background(cell, "F8FAFC")  # slate 50
            set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
            set_cell_border(cell, 
                            left=dict(val='single', sz='6', color='CBD5E1'),
                            top=dict(val='single', sz='6', color='CBD5E1'),
                            right=dict(val='single', sz='6', color='CBD5E1'),
                            bottom=dict(val='single', sz='6', color='CBD5E1'))
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(0)
            cp.paragraph_format.space_after = Pt(0)
            cp.paragraph_format.line_spacing = 1.1
            code_text = "\n".join(code_lines)
            run = cp.add_run(code_text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(51, 65, 85)

            post_p = doc.add_paragraph()
            post_p.paragraph_format.space_before = Pt(0)
            post_p.paragraph_format.space_after = Pt(4)
            continue

        # Table block: starts with |
        if line.startswith('|') and '|' in line[1:]:
            table_lines = []
            while i < total_lines and lines[i].startswith('|') and '|' in lines[i][1:]:
                table_lines.append(lines[i].rstrip('\r\n'))
                i += 1

            # Parse markdown table
            headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]
            # Filter out separator row (|---|---|)
            content_rows = []
            for t_line in table_lines[1:]:
                if re.match(r'^\s*\|?\s*[-:]+[-| :]*\|\s*$', t_line):
                    continue
                cells = [c.strip() for c in t_line.split('|')[1:-1]]
                content_rows.append(cells)

            num_cols = len(headers)
            num_rows = 1 + len(content_rows)

            table = doc.add_table(rows=num_rows, cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            # Header row styling
            hdr_cells = table.rows[0].cells
            for c_idx, h_text in enumerate(headers):
                cell = hdr_cells[c_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_background(cell, "E11D48") # Vibrant rose/pink
                set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
                set_cell_border(cell, 
                                top=dict(val='single', sz='6', color='BE123C'),
                                bottom=dict(val='single', sz='12', color='9F1239'),
                                left=dict(val='single', sz='4', color='FB7185'),
                                right=dict(val='single', sz='4', color='FB7185'))
                cp = cell.paragraphs[0]
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(0)
                run = cp.add_run(h_text)
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)

            # Content rows styling
            for r_idx, row_data in enumerate(content_rows):
                row_cells = table.rows[r_idx + 1].cells
                bg_color = "FFFFFF" if r_idx % 2 == 0 else "FDF2F4" # subtle alternate rose tint
                for c_idx in range(num_cols):
                    cell = row_cells[c_idx]
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    set_cell_background(cell, bg_color)
                    set_cell_margins(cell, top=90, bottom=90, left=110, right=110)
                    set_cell_border(cell, 
                                    top=dict(val='single', sz='4', color='F1F5F9'),
                                    bottom=dict(val='single', sz='4', color='E2E8F0'),
                                    left=dict(val='single', sz='4', color='F1F5F9'),
                                    right=dict(val='single', sz='4', color='F1F5F9'))
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_before = Pt(0)
                    cp.paragraph_format.space_after = Pt(0)
                    cell_text = row_data[c_idx] if c_idx < len(row_data) else ""
                    format_inlines(cp, cell_text, base_font="Calibri", base_size=9.5, base_color="1F2937")

            post_p = doc.add_paragraph()
            post_p.paragraph_format.space_before = Pt(2)
            post_p.paragraph_format.space_after = Pt(6)
            continue

        # Bullet list: starts with - or *
        if re.match(r'^\s*[-*]\s+', line):
            indent_level = len(line) - len(line.lstrip())
            text = re.sub(r'^\s*[-*]\s+', '', line)
            
            # Check if task checklist [ ] or [x]
            if text.startswith('[ ] '):
                bullet_char = "☐ "
                text = text[4:]
            elif text.startswith('[x] ') or text.startswith('[X] '):
                bullet_char = "☑ "
                text = text[4:]
            else:
                bullet_char = "•  "

            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25 + (indent_level // 2) * 0.2)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15

            b_run = p.add_run(bullet_char)
            b_run.bold = True
            b_run.font.name = "Arial"
            b_run.font.size = Pt(10)
            b_run.font.color.rgb = PRIMARY_COLOR

            format_inlines(p, text, base_font="Calibri", base_size=10.5, base_color="334155")
            i += 1
            continue

        # Numbered list: starts with 1. 2. etc
        m_num = re.match(r'^\s*(\d+)\.\s+(.*)', line)
        if m_num:
            num_str = m_num.group(1)
            text = m_num.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15

            n_run = p.add_run(f"{num_str}.  ")
            n_run.bold = True
            n_run.font.name = "Arial"
            n_run.font.size = Pt(10)
            n_run.font.color.rgb = SECONDARY_COLOR

            format_inlines(p, text, base_font="Calibri", base_size=10.5, base_color="334155")
            i += 1
            continue

        # Regular text paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        format_inlines(p, line, base_font="Calibri", base_size=10.5, base_color="334155")
        i += 1

    doc.save(docx_path)
    print(f"Successfully generated docx at: {docx_path}")

if __name__ == '__main__':
    md_in = sys.argv[1]
    docx_out = sys.argv[2]
    build_docx(md_in, docx_out)
